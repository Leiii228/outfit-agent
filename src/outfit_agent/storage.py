"""照片存放與 Excel 紀錄（docs/architecture.md 第 8、13 節）。"""

import asyncio
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .guards import excel_safe, excel_unescape, unique_photo_path
from .models import COLORS, NOT_OUTFIT, TYPES, Item

EXCEL_NAME = "outfits.xlsx"
PHOTOS_DIRNAME = "photos"
HEADERS = (
    "日期時間",
    "上傳者名稱",
    "上傳者 user id",
    "單品清單",
    "主要顏色",
    "風格描述",
    "這個人的穿搭總數",
    "照片存放位置",
)
COLUMN_WIDTHS = (20, 16, 22, 60, 10, 40, 16, 50)
EXCEL_RETRY_SECONDS = 10


@dataclass(frozen=True)
class Record:
    when: datetime
    name: str
    user_id: str
    items: tuple[Item, ...] | None  # None 表示「不是穿搭」
    main_color: str
    style: str
    total: int
    photo: str

    @property
    def is_outfit(self) -> bool:
        return self.items is not None


def ensure_dirs(data_dir: Path) -> None:
    (data_dir / PHOTOS_DIRNAME).mkdir(parents=True, exist_ok=True)


# ---------- 照片 ----------

def save_photo(data_dir: Path, jpeg: bytes, when: datetime, style_short: str) -> str:
    """存照片，回傳相對於 DATA_DIR 的路徑（例如 photos/2026-10-23 14-05-09 簡約休閒.jpg）。"""
    photos_dir = data_dir / PHOTOS_DIRNAME
    photos_dir.mkdir(parents=True, exist_ok=True)
    path = unique_photo_path(photos_dir, when, style_short)
    with open(path, "xb") as f:  # "x"：檔案已存在就失敗，絕不覆蓋舊照片
        f.write(jpeg)
    return f"{PHOTOS_DIRNAME}/{path.name}"


# ---------- Excel 寫入 ----------

def format_items(items: tuple[Item, ...] | None) -> str:
    if items is None:
        return NOT_OUTFIT
    return "；".join(f"{i.type}/{i.color}/{i.desc}" for i in items)


def _new_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "穿搭紀錄"
    ws.append(HEADERS)
    for idx, width in enumerate(COLUMN_WIDTHS):
        ws.column_dimensions[chr(ord("A") + idx)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def append_record(data_dir: Path, record: Record) -> None:
    """在 Excel 最後一列的下一列加入一筆紀錄（不覆蓋舊資料）。Excel 正被打開時會丟出 PermissionError。"""
    path = data_dir / EXCEL_NAME
    if not path.exists():
        _new_workbook(path)
    wb = load_workbook(path)
    ws = wb.active
    ws.append([
        record.when,
        excel_safe(record.name),
        excel_safe(record.user_id),
        excel_safe(format_items(record.items)),
        excel_safe(record.main_color),
        excel_safe(record.style),
        record.total,
        excel_safe(record.photo),
    ])
    row = ws.max_row
    ws.cell(row=row, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
    ws.cell(row=row, column=3).number_format = "@"  # user id 以文字存，避免超過 15 位數被改成 0
    wb.save(path)


async def append_record_with_retry(data_dir: Path, record: Record) -> None:
    """寫入失敗（多半是 Excel 正開著）就在終端機印警告，每 10 秒重試，直到成功為止。"""
    attempt = 0
    while True:
        try:
            await asyncio.to_thread(append_record, data_dir, record)
            if attempt:
                print(f"✓ outfits.xlsx 已寫入（重試 {attempt} 次後成功）")
            return
        except PermissionError:
            attempt += 1
            print(
                f"⚠ outfits.xlsx 可能正被 Excel 開著，{EXCEL_RETRY_SECONDS} 秒後重試（第 {attempt} 次）"
                "——關掉 Excel 就會自動存進去"
            )
            await asyncio.sleep(EXCEL_RETRY_SECONDS)


# ---------- Excel 讀取 ----------

_warned_rows: set[tuple[int, str]] = set()


def _warn_once(row_no: int, raw: str, reason: str) -> None:
    key = (row_no, raw)
    if key not in _warned_rows:
        _warned_rows.add(key)
        print(f"⚠ outfits.xlsx 第 {row_no} 列格式不對，統計時略過：{reason}")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return excel_unescape(str(value).strip())


def _parse_when(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip().lstrip("'"))
        except ValueError:
            return None
    return None


def _parse_items(text: str) -> tuple[Item, ...] | None:
    """把「類型/顏色/描述；…」轉回單品；「不是穿搭」回傳 None；格式不對丟出 ValueError。"""
    if text == NOT_OUTFIT:
        return None
    items = []
    for part in text.split("；"):
        part = part.strip()
        if not part:
            continue
        fields = part.split("/", 2)
        if len(fields) < 2:
            raise ValueError(f"單品「{part}」不是「類型/顏色/描述」格式")
        t, c = fields[0].strip(), fields[1].strip()
        desc = fields[2].strip() if len(fields) == 3 else ""
        if t not in TYPES or c not in COLORS:
            raise ValueError(f"單品「{part}」的類型或顏色不在固定選項裡")
        items.append(Item(t, c, desc))
    if not items:
        raise ValueError("單品清單是空的")
    return tuple(items)


def read_records(data_dir: Path) -> list[Record]:
    """讀出 Excel 的所有紀錄。檔案還不存在就回傳空清單；格式不對的列會印警告並略過。"""
    path = data_dir / EXCEL_NAME
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        records = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = tuple(row) + (None,) * (len(HEADERS) - len(row))
            if all(v is None for v in row[: len(HEADERS)]):
                continue
            when = _parse_when(row[0])
            user_id = _text(row[2])
            raw_items = _text(row[3])
            if when is None:
                _warn_once(row_no, str(row[0]), "日期時間看不懂")
                continue
            if not user_id.isdigit():
                _warn_once(row_no, user_id, "user id 不是數字")
                continue
            try:
                items = _parse_items(raw_items)
            except ValueError as exc:
                _warn_once(row_no, raw_items, str(exc))
                continue
            try:
                total = int(row[6] or 0)
            except (TypeError, ValueError):
                total = 0
            records.append(Record(
                when=when,
                name=_text(row[1]),
                user_id=user_id,
                items=items,
                main_color=_text(row[4]),
                style=_text(row[5]),
                total=total,
                photo=_text(row[7]),
            ))
        return records
    finally:
        wb.close()
